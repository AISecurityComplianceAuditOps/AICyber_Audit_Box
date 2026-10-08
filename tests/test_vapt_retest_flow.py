# -*- coding: utf-8 -*-
"""A VAPT session scanned, fixed, and retested -- through the real worker, API
and exports.

    pytest tests/test_vapt_retest_flow.py -v

The first scan (v1) finds five vulnerabilities on two hosts and a web
application. The client fixes some; the retest (v2) covers one of the hosts
and the application, and the pentest tracker marks a row Closed. Expected:

    still open    TLS 1.0 (10.0.0.5)          SQL injection (portal.test)
    fixed         SWEET32 (10.0.0.5, gone)    HSTS (tracker says Closed)
    new           SSH Terrapin (10.0.0.5)     Stored XSS (portal.test)
    not retested  SMB signing (10.0.0.6 was not in the retest)

The rules themselves are in test_vapt_retest.py.
"""
import io
import json
import re

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pypdf import PdfReader
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import src.api.endpoints.audit as audit_ep
import src.core.bg_worker as worker
from src.db.database import AuditReport, Base, EvidenceFile, Finding


def _nessus(items):
    hosts = {}
    for host, port, svc, pid, name, sev, cve in items:
        hosts.setdefault(host, []).append(
            f'<ReportItem port="{port}" svc_name="{svc}" protocol="tcp" severity="{sev}" pluginID="{pid}" '
            f'pluginName="{name}" pluginFamily="General">'
            + (f"<cve>{cve}</cve>" if cve else "")
            + f"<risk_factor>{['None', 'Low', 'Medium', 'High', 'Critical'][sev]}</risk_factor>"
            f"<description>{name}.</description><solution>Fix it.</solution>"
            f"<plugin_output>{name} observed.</plugin_output></ReportItem>")
    body = "".join(f'<ReportHost name="{h}"><HostProperties><tag name="host-ip">{h}</tag></HostProperties>'
                   f'{"".join(its)}</ReportHost>' for h, its in hosts.items())
    return ('<?xml version="1.0" ?><NessusClientData_v2><Report name="scan">' + body +
            "</Report></NessusClientData_v2>").encode()


V1_NESSUS = _nessus([
    ("10.0.0.5", 443, "www", 42873, "SSL Medium Strength Cipher Suites Supported (SWEET32)", 3, "CVE-2016-2183"),
    ("10.0.0.5", 443, "www", 104743, "TLS Version 1.0 Protocol Detection", 2, ""),
    ("10.0.0.6", 445, "cifs", 57608, "SMB Signing not required", 2, ""),
])
V2_NESSUS = _nessus([
    ("10.0.0.5", 443, "www", 104743, "TLS Version 1.0 Protocol Detection", 2, ""),
    ("10.0.0.5", 22, "ssh", 187315, "SSH Terrapin Prefix Truncation Weakness", 2, "CVE-2023-48795"),
])
V1_TRACKER = ("Vulnerability,Severity,Host,Description,Recommendation,Status,CVE\n"
              "SQL Injection in login form,Critical,https://portal.test/login,Injectable.,Use parameterised queries.,Open,\n"
              "Missing HSTS header,Low,https://portal.test/,HSTS not set.,Add the HSTS header.,Open,\n").encode()
V2_TRACKER = ("Vulnerability,Severity,Host,Description,Recommendation,Status,CVE\n"
              "SQL Injection in login form,Critical,https://portal.test/login,Injectable.,Use parameterised queries.,Open,\n"
              "Missing HSTS header,Low,https://portal.test/,HSTS not set.,Add the HSTS header.,Closed,\n"
              "Stored cross-site scripting in comments,High,https://portal.test/blog/comment,Stored XSS.,Encode output.,Open,\n").encode()


@pytest.fixture
def env(tmp_path, monkeypatch):
    eng = create_engine("sqlite:///" + str(tmp_path / "retest.db"))
    Base.metadata.create_all(eng)
    S = sessionmaker(bind=eng)
    monkeypatch.setattr(worker, "SessionLocal", S)
    monkeypatch.setattr(audit_ep, "SessionLocal", S)
    monkeypatch.setattr(audit_ep, "_require_auth", lambda _r: {"username": "a", "role": "admin"})
    monkeypatch.setattr(audit_ep, "_assert_session_access", lambda *a, **k: None)
    monkeypatch.setattr(audit_ep, "log_system_event", lambda *a, **k: None)
    app = FastAPI()
    app.include_router(audit_ep.router, prefix="/api")
    return {"S": S, "client": TestClient(app), "tmp": tmp_path}


def _session(S, sid, framework="VAPT Framework Controls"):
    s = S()
    rep = AuditReport(session_id=sid, session_title=sid, framework=framework, created_by="a")
    s.add(rep)
    s.commit()
    rid = rep.id
    s.close()
    return rid


def _evidence(env, rid, name, data):
    path = env["tmp"] / name
    path.write_bytes(data)
    s = env["S"]()
    ev = EvidenceFile(report_id=rid, filename=name, file_path=str(path), is_deleted=False)
    s.add(ev)
    s.commit()
    eid = ev.id
    s.close()
    return {"id": eid, "name": name, "bytes": data}


def _scan(sid, files, **kw):
    worker._run_fast_technical_vapt_bg(
        sid, [{"name": f["name"], "bytes": f["bytes"], "text": None} for f in files], set(), {},
        framework=kw.pop("framework", "VAPT Framework Controls"), ai_recommendations=False,
        evidence_files=[{"id": f["id"], "name": f["name"]} for f in files], **kw)
    with worker._bg_lock:
        return worker._bg_results.pop(sid, None)


def _rows(S, rid):
    s = S()
    try:
        return {r.control_name: r for r in s.query(Finding).filter(Finding.report_id == rid).all()}
    finally:
        s.close()


def _report(S, rid):
    s = S()
    try:
        return s.query(AuditReport).filter(AuditReport.id == rid).first()
    finally:
        s.close()


@pytest.fixture
def retested(env):
    """v1 scanned, then the retest files uploaded and scanned as a retest."""
    S = env["S"]
    rid = _session(S, "retest-flow")
    v1 = [_evidence(env, rid, "v1.nessus", V1_NESSUS), _evidence(env, rid, "v1_tracker.csv", V1_TRACKER)]
    first = _scan("retest-flow", v1)
    before = {k: (r.status, r.retest_status) for k, r in _rows(S, rid).items()}
    rounds_v1 = json.loads(_report(S, rid).vapt_rounds_json)
    v2 = [_evidence(env, rid, "v2.nessus", V2_NESSUS), _evidence(env, rid, "v2_tracker.csv", V2_TRACKER)]
    info = env["client"].get("/api/audit/vapt/retest-info?session_id=retest-flow").json()
    second = _scan("retest-flow", v2, retest=True)
    return dict(env, rid=rid, first=first, second=second, before=before, rounds_v1=rounds_v1, info=info, v1=v1, v2=v2)


SWEET32 = "SSL Medium Strength Cipher Suites Supported (SWEET32)"
TLS10 = "TLS Version 1.0 Protocol Detection"
SMB = "SMB Signing not required"
TERRAPIN = "SSH Terrapin Prefix Truncation Weakness"
SQLI = "SQL Injection in login form"
HSTS = "Missing HSTS header"
XSS = "Stored cross-site scripting in comments"


# ── the first scan is version 1 ──────────────────────────────────────────────

def test_the_first_scan_is_recorded_as_version_1(retested):
    assert retested["first"]["error"] is None
    assert set(retested["before"]) == {SWEET32, TLS10, SMB, SQLI, HSTS}
    assert all(rs is None for _, rs in retested["before"].values()), "a single scan carries no retest state"
    (v1,) = retested["rounds_v1"]
    assert v1["round"] == 1 and v1["files"] == ["v1.nessus", "v1_tracker.csv"]
    assert v1["file_ids"] == sorted(f["id"] for f in retested["v1"])
    assert v1["hosts"] == ["10.0.0.5", "10.0.0.6", "portal.test"]


def test_the_scan_workspace_is_told_which_files_are_new(retested):
    info = retested["info"]
    assert info["vapt"] is True and info["has_findings"] is True
    assert info["next_round"] == 2
    assert info["new_files"] == ["v2.nessus", "v2_tracker.csv"]


# ── the retest ───────────────────────────────────────────────────────────────

def test_each_vulnerability_gets_its_retest_status(retested):
    rows = _rows(retested["S"], retested["rid"])
    got = {k: (r.retest_status, r.status) for k, r in rows.items()}
    assert got == {
        TLS10: ("still_open", "Non-Compliant"),
        SQLI: ("still_open", "Non-Compliant"),
        SWEET32: ("fixed", "Closed"),
        HSTS: ("fixed", "Closed"),
        SMB: ("not_retested", "Non-Compliant"),
        TERRAPIN: ("new", "Non-Compliant"),
        XSS: ("new", "Non-Compliant"),
    }
    assert rows[SWEET32].final_result == "CLOSED"
    assert rows[TERRAPIN].first_round == 2 and rows[TLS10].first_round is None


def test_nothing_is_duplicated_and_v1_rows_keep_their_content(retested):
    rows = _rows(retested["S"], retested["rid"])
    assert len(rows) == 7
    assert "Injectable." in (rows[SQLI].description or "")


def test_history_says_what_each_version_found(retested):
    rows = _rows(retested["S"], retested["rid"])
    h = json.loads(rows[SWEET32].retest_history)
    assert [(e["round"], e["state"], e["status"]) for e in h] == [(1, "found", "Non-Compliant"), (2, "not_found", "Closed")]
    assert [e["state"] for e in json.loads(rows[HSTS].retest_history)] == ["found", "reported_closed"]
    assert [e["state"] for e in json.loads(rows[SMB].retest_history)] == ["found", "not_scanned"]
    assert [(e["round"], e["state"]) for e in json.loads(rows[XSS].retest_history)] == [(2, "found")]


def test_the_versions_are_recorded_with_their_counts(retested):
    rounds = json.loads(_report(retested["S"], retested["rid"]).vapt_rounds_json)
    assert [r["round"] for r in rounds] == [1, 2]
    v2 = rounds[1]
    assert v2["files"] == ["v2.nessus", "v2_tracker.csv"]
    assert v2["counts"] == {"still_open": 2, "fixed": 2, "new": 2, "not_retested": 1, "reopened": 0}
    assert "10.0.0.6" not in v2["hosts"]
    assert retested["second"]["retest"] == {"round": 2, "counts": v2["counts"]}


def test_after_the_retest_no_file_is_new(retested):
    info = retested["client"].get("/api/audit/vapt/retest-info?session_id=retest-flow").json()
    assert info["new_files"] == [] and info["next_round"] == 3


# ── the page and the reports ────────────────────────────────────────────────

def test_the_findings_api_carries_the_retest(retested):
    data = retested["client"].get("/api/audit/findings?session_id=retest-flow&include_info=true").json()
    assert [r["round"] for r in data["vapt_rounds"]] == [1, 2]
    by = {f["control_name"]: f for f in data["findings"]}
    assert by[SWEET32]["retest_status"] == "fixed" and by[SWEET32]["first_round"] == 1
    assert by[XSS]["retest_status"] == "new" and by[XSS]["first_round"] == 2
    assert [e["round"] for e in by[SWEET32]["retest_history"]] == [1, 2]


def _pdf_text(resp):
    assert resp.status_code == 200, resp.text[:300]
    return " ".join(" ".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(resp.content)).pages).split())


# The report gives the outcome, not a comparison: one line of open and closed,
# and each finding's own status Open or Closed. Reviewers found the versions
# table and the per-version grid hard to read and asked for the final status.
# Of the 7 findings: 2 fixed (closed); still open 2, new 2 and not retested 1
# are open -- not retested because nobody has confirmed a fix.
_COMPARISON_WORDS = ("Retest Summary", "Still open", "Not retested", "compared with version", "v2 retest")


def test_the_pdf_gives_the_final_open_and_closed(retested):
    text = _pdf_text(retested["client"].get("/api/audit/export/pdf?session_id=retest-flow"))
    assert "2.3.5 Retest Status" in text
    assert re.search(r"Status after the retest of \d{2} \w{3} \d{4}: 5 open, 2 closed\.", text)
    statuses = re.findall(r"Status / Scanner (Open|Closed) \| Tool", text)
    assert sorted(statuses) == ["Closed"] * 2 + ["Open"] * 5, statuses
    # The summary table says it too, one status per row.
    summary = text[text.index("2.3.4 Vulnerabilities Summary"):text.index("OVERALL SCORE")]
    assert "CVSS Score Severity Status" in summary
    rows = re.findall(r"(?:HIGH|MEDIUM|LOW|CRITICAL|INFO\w*) (Open|Closed)\b", summary)
    assert sorted(rows) == ["Closed"] * 2 + ["Open"] * 5, rows
    for word in _COMPARISON_WORDS:
        assert word not in text, word


def test_the_first_version_exports_as_it_was(retested):
    text = _pdf_text(retested["client"].get("/api/audit/export/pdf?session_id=retest-flow&version=1"))
    assert "Retest Status" not in text
    assert "CVSS Score Severity Status" not in text, "a version scanned once keeps its table"
    assert "Terrapin" not in text and "Stored cross-site scripting" not in text
    assert "SWEET32" in text


def test_the_status_line_can_be_left_out(retested):
    text = _pdf_text(retested["client"].get("/api/audit/export/pdf?session_id=retest-flow&retest_summary=false"))
    assert "Retest Status" not in text and "Terrapin" in text


def test_the_word_report_gives_each_finding_open_or_closed(retested):
    r = retested["client"].get("/api/audit/export/docx?session_id=retest-flow")
    assert r.status_code == 200, r.text[:300]
    d = Document(io.BytesIO(r.content))
    paras = [p.text for p in d.paragraphs]
    assert "2.3.5 Retest Status" in paras
    assert any(re.search(r"Status after the retest of .*: 5 open, 2 closed\.", p) for p in paras)
    statuses = [m.group(1) for p in paras for m in [re.search(r"\|\s+Status: (Open|Closed)\b", p)] if m]
    assert sorted(statuses) == ["Closed"] * 2 + ["Open"] * 5, statuses
    summary = next(t for t in d.tables if [c.text for c in t.rows[0].cells][:2] == ["Sr. No.", "Vulnerabilities"])
    assert [c.text for c in summary.rows[0].cells] == ["Sr. No.", "Vulnerabilities", "CVSS Score", "Severity", "Status"]
    assert sorted(r.cells[4].text for r in summary.rows[1:-1]) == ["Closed"] * 2 + ["Open"] * 5
    everything = " ".join(paras + [c.text for t in d.tables for row in t.rows for c in row.cells])
    for word in _COMPARISON_WORDS:
        assert word not in everything, word


# ── what must not change ─────────────────────────────────────────────────────

def test_a_normal_rerun_makes_the_session_one_version_again(retested):
    S, rid = retested["S"], retested["rid"]
    _scan("retest-flow", retested["v1"] + retested["v2"])
    rounds = json.loads(_report(S, rid).vapt_rounds_json)
    assert [r["round"] for r in rounds] == [1]
    assert sorted(rounds[0]["files"]) == ["v1.nessus", "v1_tracker.csv", "v2.nessus", "v2_tracker.csv"]
    assert all(r.retest_status is None and r.retest_history is None for r in _rows(S, rid).values())


def test_pqc_is_untouched(env):
    rid = _session(env["S"], "pqc-flow", framework="PQC Framework Controls")
    f = _evidence(env, rid, "scan.nessus", V1_NESSUS)
    _scan("pqc-flow", [f], framework="PQC Framework Controls", retest=True)
    assert _report(env["S"], rid).vapt_rounds_json is None
    assert all(r.retest_status is None for r in _rows(env["S"], rid).values())


def test_a_non_vapt_session_reports_no_retest(env):
    _session(env["S"], "iso-flow", framework="ISO 27001")
    assert env["client"].get("/api/audit/vapt/retest-info?session_id=iso-flow").json() == {"vapt": False}


def _start(client, sid, framework):
    return client.post("/api/audit/start", json={
        "session_id": sid, "selected_sls": [1], "model_choice": "x",
        "audit_mode": "Technical findings only", "current_framework": framework,
        "vapt_retest": True, "ai_recommendations": False})


def test_a_retest_is_refused_where_it_cannot_apply(env):
    client = env["client"]
    _session(env["S"], "iso-start", framework="ISO 27001")
    r = _start(client, "iso-start", "ISO 27001")
    assert r.status_code == 400 and "not a VAPT scanner session" in r.json()["detail"]

    rid = _session(env["S"], "empty-start")
    _evidence(env, rid, "v1.nessus", V1_NESSUS)
    r = _start(client, "empty-start", "VAPT Framework Controls")
    assert r.status_code == 400 and "no earlier scan" in r.json()["detail"]


def test_a_retest_with_no_new_file_is_refused(retested):
    r = _start(retested["client"], "retest-flow", "VAPT Framework Controls")
    assert r.status_code == 400 and "already scanned" in r.json()["detail"]
    # Refusing changed nothing.
    assert len(_rows(retested["S"], retested["rid"])) == 7
