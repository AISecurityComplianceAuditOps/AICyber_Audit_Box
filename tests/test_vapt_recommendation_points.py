# -*- coding: utf-8 -*-
"""A VAPT recommendation in the reports is points, as on the finding's card.

    pytest tests/test_vapt_recommendation_points.py -v

The card has long shown a recommendation as points (app.js
formatRemediationSteps): numbered steps as a numbered list, the AI's "- " lines
as bullets, prose one bullet per sentence. The PDF and Word reports printed it
as one block -- and advice read out of a scanner's PDF kept that PDF's line
wrapping, so a delivered report's lines stopped mid-sentence:

    ... the application specifies the structure of the query, leaving placeholders for each
    item of user input; second, ...

Both reports now split it as the card does, with the wrapping joined.
"""
import io
import re

from src.core.report_exporter import (_VAPT_CLOSED_NOTE, _vapt_points, export_docx_report,
                                      export_pdf_report)

# Burp's own advice for SQL injection, as read out of PortSwigger's PDF report:
# hard line breaks mid-sentence.
BURP_SQLI = """Parameterized queries (also known as prepared statements) for all database access. This method
uses two steps to incorporate potentially tainted data into SQL queries: first, the application specifies the structure of the query, leaving placeholders for each
item of user input; second, the application specifies the contents of each placeholder. Because the structure of the query has already been defined in the first
step, it is not possible for malformed data in the second step to interfere with the query structure. You should review the documentation for your database and
application platform, e.g. MySQL, to determine the appropriate APIs. Upgrade to 2.4.51 or later."""


def test_prose_is_one_point_per_sentence_with_the_wrapping_joined():
    pre, items, numbered = _vapt_points(BURP_SQLI)
    assert (pre, numbered, len(items)) == ("", False, 5)
    assert all("\n" not in i for i in items)
    assert "leaving placeholders for each item of user input" in items[1]
    assert items[3].endswith("e.g. MySQL, to determine the appropriate APIs.")   # not split at "e.g."
    assert items[4] == "Upgrade to 2.4.51 or later."                            # nor inside a version


def test_the_ai_points_and_steps_stay_as_written():
    assert _vapt_points("- Disable TLS 1.0 on 10.0.0.5:443.\n- Restart nginx.") == \
        ("", ["Disable TLS 1.0 on 10.0.0.5:443.", "Restart nginx."], False)
    assert _vapt_points("1. Edit `ssl_protocols`.\n2. Reload nginx.\n3. Re-test with testssl.sh.") == \
        ("", ["Edit `ssl_protocols`.", "Reload nginx.", "Re-test with testssl.sh."], True)


def test_steps_after_a_heading_are_numbered_under_it():
    pre, items, numbered = _vapt_points(
        "General guidance from MITRE CWE-319, not from the report: 1. Encrypt it. 2. Use TLS 1.2 or later.")
    assert pre == "General guidance from MITRE CWE-319, not from the report:"
    assert (items, numbered) == (["Encrypt it.", "Use TLS 1.2 or later."], True)


def test_one_sentence_or_one_step_stays_a_paragraph():
    assert _vapt_points("Use TLS.") == ("", ["Use TLS."], False)
    assert _vapt_points("1. Replace string concatenation with bound parameters.") == \
        ("", ["1. Replace string concatenation with bound parameters."], False)


def _findings():
    return [dict(control_id="VAPT-4", title="SQL injection", severity="HIGH", status="Non-Compliant",
                 final_result="NON_COMPLIANT", description="d", target="https://shop.test/catalog/filter",
                 evidence_snippet="proof", recommendation=BURP_SQLI,
                 remediation_actionable="1. Use bound parameters.\n2. Re-test the category parameter."),
            dict(control_id="VAPT-5", title="Cleartext Transmission", severity="LOW", status="Closed",
                 final_result="CLOSED", description="d", target="https://shop.test/", evidence_snippet="p",
                 recommendation="Use TLS for the API. Redirect HTTP to HTTPS.")]


def test_the_word_report_gives_points():
    from docx import Document
    d = Document(io.BytesIO(export_docx_report("t", _findings(), [], "FINAL", audit_type="vapt")))
    paras = [p.text for p in d.paragraphs]
    bullets = [p for p in paras if p.startswith("•\t")]
    assert len(bullets) == 5 + 2, bullets                       # the open one's 5, the closed one's 2
    assert any("leaving placeholders for each item of user input" in b for b in bullets)
    assert "1.\tUse bound parameters." in paras and "2.\tRe-test the category parameter." in paras
    # A closed finding: the note stays one paragraph, its own advice is the points.
    assert _VAPT_CLOSED_NOTE in paras
    i = paras.index("Original recommendation (for reference):")
    assert paras[i + 1:i + 3] == ["•\tUse TLS for the API.", "•\tRedirect HTTP to HTTPS."]


def test_the_pdf_writes_the_points_left_aligned_and_whole():
    from pypdf import PdfReader
    out = export_pdf_report("t", _findings(), [], "FINAL", audit_type="vapt")
    text = "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(out)).pages)
    flat = re.sub(r"\s+", " ", text)
    assert "Parameterized queries (also known as prepared statements) for all database access." in flat
    assert "1. Use bound parameters. 2. Re-test the category parameter." in flat
    assert "Remediation Status: No action required - this finding is closed (remediated)." in flat
    assert "Original recommendation (for reference): Use TLS for the API. Redirect HTTP to HTTPS." in flat
