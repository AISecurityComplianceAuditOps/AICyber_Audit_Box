# -*- coding: utf-8 -*-
"""A finding the pentest report records as closed is counted as Closed.

It was stored as Informational (severity rewritten to INFO), so a VAPT session
listed a remediated High among the scanner's informational notes, and the
reports counted it there. Now the finding keeps the severity the report gave
it, its status is "Closed", and the page and both reports count it under
Closed only -- in the order Critical, High, Medium, Low, Informational,
Closed, Total. Closed is a VAPT idea: no other framework's counters change.

The page's own functions are run under Node where it is available.
"""
import io
import json
import os
import re
import shutil
import subprocess

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import src.core.bg_worker as worker
from src.db.database import AuditReport, Base, Finding

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
APP_JS = os.path.join(ROOT, "src", "api", "static", "app.js")
INDEX = os.path.join(ROOT, "src", "api", "static", "index.html")

TRACKER = ("Vulnerability,Severity,Host,Description,Recommendation,Status,CVE\n"
           "SQL Injection in login form,Critical,https://portal.test/login,Injectable.,Use parameterised queries.,Open,\n"
           "Missing HSTS header,Low,https://portal.test/,HSTS not set.,Add the HSTS header.,Closed,\n"
           "Weak TLS ciphers,High,https://portal.test/,3DES offered.,Disable 3DES.,Fixed,\n")


# -- the worker ---------------------------------------------------------------

def test_the_worker_stores_a_closed_finding_as_closed_with_its_severity(tmp_path, monkeypatch):
    eng = create_engine("sqlite:///" + str(tmp_path / "vapt.db"))
    Base.metadata.create_all(eng)
    S = sessionmaker(bind=eng)
    monkeypatch.setattr(worker, "SessionLocal", S)
    s = S()
    s.add(AuditReport(session_id="c1", session_title="c1", framework="VAPT"))
    s.commit()
    s.close()
    worker._run_fast_technical_vapt_bg("c1", [{"name": "tracker.csv", "bytes": TRACKER.encode(), "text": None}],
                                       selected_sls=[], framework="VAPT")
    s = S()
    try:
        rep = s.query(AuditReport).filter(AuditReport.session_id == "c1").first()
        got = {r.control_name: (r.status, r.severity, r.final_result)
               for r in s.query(Finding).filter(Finding.report_id == rep.id)}
    finally:
        s.close()
    by = {t: next(v for k, v in got.items() if t in k) for t in ("SQL Injection", "HSTS", "TLS")}
    assert by["SQL Injection"] == ("Non-Compliant", "CRITICAL", None)
    # The verdict CLOSED too, so an Accept (which keeps the verdict) leaves it closed.
    assert by["HSTS"] == ("Closed", "LOW", "CLOSED")
    assert by["TLS"] == ("Closed", "HIGH", "CLOSED")


# -- the reports --------------------------------------------------------------

_CLOSED_NOTE = "Status in report: Closed / remediated\n| row |"


def _findings():
    base = {"description": "d", "recommendation": "r", "target": "https://portal.test/"}
    return [
        dict(base, control_id="VAPT-1", title="SQL Injection", severity="HIGH", status="Non-Compliant",
             final_result="NON_COMPLIANT", evidence_snippet="proof"),
        dict(base, control_id="VAPT-2", title="Missing HSTS", severity="LOW", status="Closed",
             final_result="CLOSED", evidence_snippet=_CLOSED_NOTE),
        dict(base, control_id="VAPT-3", title="Server banner", severity="INFO", status="Informational",
             final_result="NON_COMPLIANT", evidence_snippet="Server: nginx"),
        # Accept confirms a closed finding (keeps its verdict CLOSED): still closed.
        dict(base, control_id="VAPT-4", title="Weak TLS", severity="MEDIUM", status="Accepted",
             final_result="CLOSED", evidence_snippet=_CLOSED_NOTE),
        # Accept on an open one: still an open Medium. So is one the auditor
        # reopened, although its proof still carries the report's closed line.
        dict(base, control_id="VAPT-5", title="Clickjacking", severity="MEDIUM", status="Accepted",
             final_result="NON_COMPLIANT", evidence_snippet=_CLOSED_NOTE),
        dict(base, control_id="VAPT-6", title="Rejected one", severity="CRITICAL", status="Rejected",
             final_result="NON_COMPLIANT", evidence_snippet="x"),
    ]


HEADER = ["Critical", "High", "Medium", "Low", "Informational", "Closed", "Total Findings"]
COUNTS = ["0", "1", "1", "0", "1", "2", "5"]
SENTENCE = ("Based on the assessment, 2 open vulnerabilities have been found in the target scope which are "
            "categorized as follows, with 1 further informational observation(s) and 2 finding(s) the report "
            "records as closed:")


def test_the_vapt_docx_counts_closed_apart_in_the_mentors_order():
    from docx import Document
    from src.core.report_exporter import export_docx_report
    out = export_docx_report("closed check", _findings(), [], "FINAL", audit_type="vapt")
    d = Document(io.BytesIO(out))
    tables = [[[c.text.strip() for c in row.cells] for row in t.rows] for t in d.tables]
    summary = [t for t in tables if t and t[0] == HEADER]
    assert summary and summary[0][1] == COUNTS, [t[:2] for t in tables[:12]]
    assert SENTENCE in [p.text for p in d.paragraphs]


def test_the_vapt_pdf_counts_closed_apart_in_the_mentors_order():
    from pypdf import PdfReader
    from src.core.report_exporter import export_pdf_report
    out = export_pdf_report("closed check", _findings(), [], "FINAL", audit_type="vapt")
    text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(out)).pages)
    flat = re.sub(r"\s+", " ", text)
    assert " ".join(HEADER) + " " + " ".join(COUNTS) in flat
    assert "2.3.1 Findings Overview: " + SENTENCE in flat
    assert flat.count("Closed (recorded as remediated in the report)") == 2


@pytest.mark.parametrize("info,closed,tail", [
    (0, 0, "as follows:"),
    (3, 0, "as follows, with 3 further informational observation(s):"),
    (0, 2, "as follows, with 2 finding(s) the report records as closed:"),
])
def test_the_overview_sentence_reads_right_with_any_counts(info, closed, tail):
    from src.core.report_exporter import _vapt_overview_sentence
    assert _vapt_overview_sentence(4, info, closed).endswith(tail)


# -- the page -----------------------------------------------------------------

def _src():
    return io.open(APP_JS, encoding="utf-8").read()


def _function(src, name):
    start = src.index(f"function {name}(")
    depth, i = 0, src.index("{", start)
    while True:
        ch = src[i]
        depth += ch == "{"
        depth -= ch == "}"
        i += 1
        if depth == 0:
            return src[start:i]


PAGE_FINDINGS = [
    {"severity": "HIGH", "status": "Non-Compliant", "final_result": "NON_COMPLIANT"},
    {"severity": "HIGH", "status": "Non-Compliant", "final_result": "NON_COMPLIANT"},
    {"severity": "LOW", "status": "Non-Compliant", "final_result": "NON_COMPLIANT"},
    {"severity": "LOW", "status": "Closed", "final_result": "CLOSED"},
    {"severity": "MEDIUM", "status": "Accepted", "final_result": "CLOSED",
     "evidence_snippet": "Status in report: Closed / remediated"},
    {"severity": "INFO", "status": "Informational", "final_result": "NON_COMPLIANT"},
    {"severity": "CRITICAL", "status": "Rejected", "final_result": "NON_COMPLIANT"},
]


def _stats(framework):
    node = shutil.which("node")
    if not node:
        pytest.skip("node not available")
    src = _src()
    script = "\n".join([
        "const els = {};",
        "const document = { getElementById: id => (els[id] = els[id] || { innerText: '' }) };",
        f"let findingsSessionFramework = {json.dumps(framework)};",
        f"let findingsList = {json.dumps(PAGE_FINDINGS)};",
        # the verdict branch of isFindingCompliant (every finding here has one)
        "function isFindingCompliant(f) { return String(f.final_result || '').toUpperCase() === 'COMPLIANT'; }",
    ] + [_function(src, n) for n in ("isVaptOnlySession", "isFindingClosed", "isFindingInformational",
                                     "severityBand", "calculateSeverityStats")])
    script += "\ncalculateSeverityStats();\nconst out = {};\n"
    script += "for (const k in els) out[k] = String(els[k].innerText);\n"
    script += "process.stdout.write(JSON.stringify(out));\n"
    out = subprocess.run([node, "-e", script], capture_output=True, text=True, timeout=30)
    assert out.returncode == 0, out.stderr
    return json.loads(out.stdout)


def test_a_vapt_session_counts_closed_apart():
    got = _stats("VAPT")
    assert [got["count-vapt-" + k] for k in ("critical", "high", "medium", "low", "info", "closed", "total")] \
        == ["0", "2", "0", "1", "1", "2", "6"]


def test_no_other_framework_has_closed_findings():
    """ISO keeps every count it had: the closed-status rows count by severity."""
    got = _stats("ISO 27001")
    assert got["count-vapt-closed"] == "0"
    assert (got["count-noncompliant"], got["count-p2"], got["count-p3"], got["count-p4"]) == ("5", "2", "1", "2")
    assert got["count-compliant"] == "0"


def test_the_vapt_row_is_in_the_mentors_order():
    html = io.open(INDEX, encoding="utf-8").read()
    row = html[html.index('id="kpi-row-vapt"'):]
    row = row[:row.index('id="count-vapt-total"')]
    labels = re.findall(r'<span class="kpi-label">([^<]+)</span>', row)
    assert labels == ["Critical", "High", "Medium", "Low", "Informational", "Closed", "Total"]
    assert 'id="kpi-row-standard"' in html                  # the other frameworks' row is still there


def test_the_filters_keep_closed_out_of_the_open_views():
    src = _src()
    chain = src[src.index('if (valLower !== "all") {'):][:3500]
    assert 'list = list.filter(f => isFindingInformational(f) && !isFindingClosed(f));' in chain
    assert 'valLower === "closed"' in chain
    assert "severityBand(f.severity) === wanted && !isFindingClosed(f)" in chain
