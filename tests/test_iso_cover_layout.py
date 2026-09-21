# -*- coding: utf-8 -*-
"""The ISO cover's lower block is one bordered box, not two.

    pytest tests/test_iso_cover_layout.py -v

WHY THIS EXISTS

Word draws a run of consecutive paragraphs that share the same border as a
single box. The cover's lower block -- the auditor logo, "Audit Conducted By:",
the name, the firm, the address and the contact lines -- is one such run.

The auditor's logo is inserted into that run at export time with
insert_paragraph_before(), which creates a BARE paragraph. A paragraph with no
border in the middle of a bordered run ends the first box and starts a second,
so the delivered cover showed two stacked boxes with the shield stranded in the
gap between them, overlapping the lower box's top edge.

Nothing raised, and the DOCX and the PDF agreed with each other -- they were
both wrong in the same way, because the PDF is a LibreOffice render of this very
DOCX. It could only be seen by looking at the page.
"""
import io

import pytest

from docx import Document

from src.core.report_exporter import _export_iso_template_docx


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

FINDING = {
    "control_id": "8.17", "control_name": "Clock Synchronization",
    "severity": "HIGH", "status": "Non-Compliant",
    "final_result": "NON_COMPLIANT", "description": "d",
    "recommendation": "r", "evidence_snippet": "e",
}


def _cover_paragraphs():
    buf = _export_iso_template_docx(
        session_title="ISO 27001 - Test - 22 Sept 2026",
        findings=[dict(FINDING)], resolved_list=[], status="Final", comments="",
        metadata={"brand_firm": "Dhiware Technologies Pvt Ltd"})
    data = buf.getvalue() if hasattr(buf, "getvalue") else buf
    if not data:
        pytest.skip("the ISO template did not render on this machine")
    doc = Document(io.BytesIO(data))
    end = next((i for i, p in enumerate(doc.paragraphs)
                if p.text.strip() == "Contents"), len(doc.paragraphs))
    return doc.paragraphs[:end]


def _has_border(p):
    pr = p._p.find(W + "pPr")
    return pr is not None and pr.find(W + "pBdr") is not None


def _has_picture(p):
    return bool(p._p.findall(".//" + A + "blip"))


def test_the_logo_paragraph_carries_the_block_border():
    """The defect, stated directly: a bare paragraph split the box in two."""
    cover = _cover_paragraphs()
    logo = [p for p in cover if _has_picture(p)]
    assert logo, "no logo on the cover at all"
    # The auditee slot at the top may be an empty placeholder; the auditor's
    # logo is the one inside the lower bordered run.
    in_block = [p for p in logo if _has_border(p)]
    assert in_block, (
        "the auditor logo sits in a borderless paragraph, which ends the "
        "bordered box above it and starts a new one below -- the logo renders "
        "stranded in the gap between two boxes")


def test_the_lower_block_is_one_unbroken_run():
    """From the logo to the email line, every paragraph shares the border.

    Checked as a run rather than per paragraph: one borderless paragraph
    anywhere inside it is what breaks the box, wherever it lands.
    """
    cover = _cover_paragraphs()
    start = next((i for i, p in enumerate(cover)
                  if p.text.strip().startswith("Audit Conducted By")), None)
    assert start is not None, "the cover has no 'Audit Conducted By:' line"

    # Walk backwards over the logo paragraph(s) that precede it.
    while start > 0 and (_has_picture(cover[start - 1]) or _has_border(cover[start - 1])):
        start -= 1

    end = next((i for i, p in enumerate(cover)
                if p.text.strip().startswith("Email")), None)
    assert end is not None, "the cover has no email line"

    unbordered = [i for i in range(start, end + 1) if not _has_border(cover[i])]
    assert not unbordered, (
        "paragraph(s) %r inside the cover's lower block carry no border, so "
        "Word will draw it as more than one box" % (unbordered,))


def test_the_logo_comes_before_the_conducted_by_line():
    cover = _cover_paragraphs()
    logo_idx = next((i for i, p in enumerate(cover)
                     if _has_picture(p) and _has_border(p)), None)
    conducted = next((i for i, p in enumerate(cover)
                      if p.text.strip().startswith("Audit Conducted By")), None)
    assert logo_idx is not None and conducted is not None
    assert logo_idx < conducted, "the auditor logo is below 'Audit Conducted By:'"
